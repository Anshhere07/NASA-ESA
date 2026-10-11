"""
Unit tests for exporter.py and workflow.py

Tests:
1. Configuration loading and validation.
2. Overwrite policies ('overwrite', 'error', 'skip').
3. CSV file creation and structure.
4. Excel workbook styling, precision, sheet names, frozen panes, and auto-filters.
5. Text summary contents, extrema, and scientific provenance statements.
6. End-to-end workflow execution.
"""

import os
import unittest
from pathlib import Path
import tempfile
import json

import pandas as pd
import openpyxl
import spiceypy as spice

from workflow import load_and_validate_config, run_workflow
from exporter import (
    check_overwrite_policy,
    export_to_csv,
    export_to_excel,
    generate_text_summary,
    export_text_summary,
)


class TestConfigValidation(unittest.TestCase):
    """Test configuration validation."""

    def setUp(self):
        self.valid_config = {
            "target": "4",
            "observer": "399",
            "reference_frame": "J2000",
            "aberration_correction": "LT+S",
            "start_utc": "2026-10-11T00:00:00",
            "end_utc": "2026-10-12T00:00:00",
            "sampling_interval_hours": 1.0,
            "output_formats": ["csv", "xlsx", "txt"],
            "kernel_directory": "kernels",
            "output_directory": "output",
            "overwrite_policy": "overwrite",
        }

    def test_valid_config(self):
        cfg = load_and_validate_config(self.valid_config)
        self.assertEqual(cfg["target"], "4")
        self.assertEqual(cfg["overwrite_policy"], "overwrite")

    def test_missing_mandatory_keys(self):
        invalid = dict(self.valid_config)
        del invalid["target"]
        with self.assertRaises(ValueError) as ctx:
            load_and_validate_config(invalid)
        self.assertIn("missing mandatory fields", str(ctx.exception))

    def test_invalid_interval(self):
        invalid = dict(self.valid_config)
        invalid["sampling_interval_hours"] = -2.0
        with self.assertRaises(ValueError) as ctx:
            load_and_validate_config(invalid)
        self.assertIn("sampling_interval_hours", str(ctx.exception))

    def test_unsupported_output_format(self):
        invalid = dict(self.valid_config)
        invalid["output_formats"] = ["csv", "pdf"]
        with self.assertRaises(ValueError) as ctx:
            load_and_validate_config(invalid)
        self.assertIn("Unsupported output format", str(ctx.exception))

    def test_invalid_overwrite_policy(self):
        invalid = dict(self.valid_config)
        invalid["overwrite_policy"] = "destroy_everything"
        with self.assertRaises(ValueError) as ctx:
            load_and_validate_config(invalid)
        self.assertIn("Invalid 'overwrite_policy'", str(ctx.exception))

    def test_nonexistent_kernel_directory(self):
        invalid = dict(self.valid_config)
        invalid["kernel_directory"] = "nonexistent_kernels_dir_abc"
        with self.assertRaises(FileNotFoundError):
            load_and_validate_config(invalid)


class TestExporter(unittest.TestCase):
    """Test CSV, Excel, and Text Summary exports."""

    def setUp(self):
        self.tmp_dir = tempfile.TemporaryDirectory()
        self.out_dir = Path(self.tmp_dir.name)

        self.sample_data = pd.DataFrame([
            {
                "UTC": "2026-10-11T00:00:00",
                "ET_seconds": 844948869.182,
                "Target": "4",
                "Observer": "399",
                "Frame": "J2000",
                "X_km": -143655422.37,
                "Y_km": 172712190.22,
                "Z_km": 80276150.31,
                "VX_km_s": -13.9558,
                "VY_km_s": -24.4560,
                "VZ_km_s": -9.9345,
            },
            {
                "UTC": "2026-10-11T06:00:00",
                "ET_seconds": 844970469.182,
                "Target": "4",
                "Observer": "399",
                "Frame": "J2000",
                "X_km": -143955518.23,
                "Y_km": 172183814.77,
                "Z_km": 80061501.95,
                "VX_km_s": -13.8323,
                "VY_km_s": -24.4676,
                "VZ_km_s": -9.9406,
            },
        ])

        self.config = {
            "target": "4",
            "target_name": "MARS BARYCENTER",
            "observer": "399",
            "observer_name": "EARTH",
            "reference_frame": "J2000",
            "aberration_correction": "LT+S",
            "start_utc": "2026-10-11T00:00:00",
            "end_utc": "2026-10-11T06:00:00",
            "sampling_interval_hours": 6.0,
            "kernel_directory": "kernels",
        }

    def tearDown(self):
        self.tmp_dir.cleanup()

    def test_overwrite_policy_error(self):
        file_path = self.out_dir / "test_file.csv"
        file_path.write_text("existing content", encoding="utf-8")

        with self.assertRaises(FileExistsError):
            check_overwrite_policy(file_path, policy="error")

        self.assertFalse(check_overwrite_policy(file_path, policy="skip"))
        self.assertTrue(check_overwrite_policy(file_path, policy="overwrite"))

    def test_csv_export(self):
        csv_file = self.out_dir / "output.csv"
        out = export_to_csv(self.sample_data, csv_file, overwrite_policy="overwrite")
        self.assertTrue(csv_file.exists())

        loaded_df = pd.read_csv(csv_file)
        self.assertEqual(len(loaded_df), 2)
        self.assertEqual(list(loaded_df.columns), list(self.sample_data.columns))

    def test_excel_export_formatting(self):
        xlsx_file = self.out_dir / "output.xlsx"
        export_to_excel(
            self.sample_data,
            xlsx_file,
            sheet_name="SPICE_Trajectory",
            overwrite_policy="overwrite",
        )
        self.assertTrue(xlsx_file.exists())

        wb = openpyxl.load_workbook(xlsx_file)
        self.assertIn("SPICE_Trajectory", wb.sheetnames)
        ws = wb["SPICE_Trajectory"]

        # Check frozen header
        self.assertEqual(ws.freeze_panes, "A2")

        # Check auto filter exists
        self.assertIsNotNone(ws.auto_filter.ref)

        # Check precision formatting on row 2
        # Column 2: ET_seconds (#,##0.000)
        self.assertEqual(ws.cell(row=2, column=2).number_format, "#,##0.000")
        # Column 6: X_km (#,##0.00)
        self.assertEqual(ws.cell(row=2, column=6).number_format, "#,##0.00")
        # Column 9: VX_km_s (0.0000)
        self.assertEqual(ws.cell(row=2, column=9).number_format, "0.0000")

    def test_text_summary_content(self):
        summary = generate_text_summary(self.sample_data, self.config)

        self.assertIn("MARS BARYCENTER", summary)
        self.assertIn("EARTH", summary)
        self.assertIn("J2000", summary)
        self.assertIn("LT+S", summary)
        self.assertIn("Position Units:           kilometres (km)", summary)
        self.assertIn("Velocity Units:           kilometres per second (km/s)", summary)
        self.assertIn("[First Successful Record]", summary)
        self.assertIn("[Last Successful Record]", summary)
        self.assertIn("X Component (km):", summary)
        self.assertIn("Range Distance (km):", summary)

        # Scientific provenance statement
        self.assertIn("SCIENTIFIC DATA PROVENANCE NOTICE", summary)
        self.assertIn("mathematically", summary)
        self.assertIn("NOT direct in-situ spacecraft", summary)

        # Save summary
        txt_path = self.out_dir / "summary.txt"
        export_text_summary(summary, txt_path)
        self.assertTrue(txt_path.exists())
        self.assertEqual(txt_path.read_text(encoding="utf-8"), summary)


class TestWorkflowEndToEnd(unittest.TestCase):
    """Test full workflow execution with config.json."""

    def test_run_workflow(self):
        spice.kclear()
        result = run_workflow("config.json")

        self.assertIn("dataframe", result)
        self.assertIn("exported_files", result)
        self.assertGreater(result["success_count"], 0)

        exported = result["exported_files"]
        self.assertIn("csv", exported)
        self.assertIn("xlsx", exported)
        self.assertIn("txt", exported)

        self.assertTrue(exported["csv"].exists())
        self.assertTrue(exported["xlsx"].exists())
        self.assertTrue(exported["txt"].exists())

        # Reopen CSV
        df_csv = pd.read_csv(exported["csv"])
        self.assertEqual(len(df_csv), 25)  # 24h interval with 1h step -> 25 points

        # Reopen Excel
        df_excel = pd.read_excel(exported["xlsx"], sheet_name="SPICE_Trajectory")
        self.assertEqual(len(df_excel), 25)

        spice.kclear()


if __name__ == "__main__":
    unittest.main()
