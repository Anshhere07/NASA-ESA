"""
Comprehensive Integration Test Suite for SPICE Trajectory Extraction Pipeline.

Systematically verifies all 12 operational capabilities:
1. Package imports.
2. Configuration loading and validation.
3. Kernel discovery and validation.
4. UTC-to-ET conversion.
5. Target and observer verification.
6. Ephemeris coverage checks.
7. Real 6-component state vector extraction.
8. Time-series sampling.
9. CSV export and verification.
10. Excel export, styling, number formats, and frozen panes.
11. Summary report generation and scientific provenance notice.
12. Error handling and guaranteed kernel cleanup.
"""

import os
import unittest
from pathlib import Path
import tempfile
import json

import numpy as np
import pandas as pd
import openpyxl
import spiceypy as spice

# 1. Package & Module Imports Verification
import spice_manager
from spice_manager import (
    SpiceManager,
    utc_to_et,
    is_lsk_loaded,
    KernelNotFoundError,
    MissingKernelError,
    UnsupportedTargetError,
    CoverageError,
    SpiceTimeConversionError,
)
import trajectory
from trajectory import (
    get_state_vector,
    extract_time_series,
    validate_extraction_parameters,
    StateVector,
    TrajectoryExtractionResult,
    InvalidParameterError,
    EphemerisExtractionError,
    DATAFRAME_COLUMNS,
    FAILURE_COLUMNS,
)
import exporter
from exporter import (
    export_to_csv,
    export_to_excel,
    generate_text_summary,
    export_text_summary,
    check_overwrite_policy,
)
import workflow
from workflow import load_and_validate_config, run_workflow
import main
from main import run_pipeline_with_cleanup, validate_dependencies_cli


class TestCompleteApplicationPipeline(unittest.TestCase):
    """Full 12-point integration test suite."""

    def setUp(self):
        spice.kclear()
        self.kernels_dir = Path("kernels")
        self.lsk = self.kernels_dir / "naif0012.tls"
        self.pck = self.kernels_dir / "pck00010.tpc"
        self.spk = self.kernels_dir / "de440s.bsp"
        self.has_real_kernels = self.lsk.exists() and self.spk.exists() and self.pck.exists()

    def tearDown(self):
        spice.kclear()

    # Point 1: Package Imports
    def test_01_package_imports(self):
        self.assertTrue(hasattr(spice, "tkvrsn"))
        self.assertTrue(hasattr(np, "array"))
        self.assertTrue(hasattr(pd, "DataFrame"))
        self.assertTrue(hasattr(openpyxl, "Workbook"))
        self.assertIsNotNone(spice.tkvrsn("TOOLKIT"))

    # Point 2: Configuration loading and validation
    def test_02_config_loading_and_validation(self):
        config = load_and_validate_config("config.json")
        self.assertIn("target", config)
        self.assertIn("observer", config)
        self.assertGreater(config["sampling_interval_hours"], 0)
        self.assertIn(config["overwrite_policy"], ["overwrite", "error", "skip"])

        # Reject invalid sampling intervals
        bad_config = dict(config)
        bad_config["sampling_interval_hours"] = -5.0
        with self.assertRaises(ValueError):
            load_and_validate_config(bad_config)

    # Point 3: Kernel discovery and validation
    def test_03_kernel_discovery_and_validation(self):
        if not self.has_real_kernels:
            self.skipTest("Real kernels unavailable for discovery test.")
        with SpiceManager(self.kernels_dir) as sm:
            discovered = sm.discover_kernels()
            self.assertGreater(len(discovered["lsk"]), 0)
            self.assertGreater(len(discovered["spk"]), 0)
            self.assertGreater(len(discovered["pck"]), 0)

            valid_lsk = sm.validate_kernel_file(self.lsk)
            self.assertTrue(valid_lsk.exists())

            with self.assertRaises(KernelNotFoundError):
                sm.validate_kernel_file("kernels/non_existent.bsp")

    # Point 4: UTC-to-ET conversion
    def test_04_utc_to_et_conversion(self):
        if not self.has_real_kernels:
            self.skipTest("Real kernels unavailable for time conversion.")
        with SpiceManager(self.kernels_dir) as sm:
            # Without LSK -> MissingKernelError
            self.assertFalse(is_lsk_loaded())
            with self.assertRaises(MissingKernelError):
                utc_to_et("2026-10-11T00:00:00")

            # Load LSK and convert
            sm.load_kernel(self.lsk)
            et = utc_to_et("2026-10-11T00:00:00")
            self.assertIsInstance(et, float)
            self.assertAlmostEqual(et, 844948869.182, delta=0.01)

            # Rejects invalid syntax
            with self.assertRaises(SpiceTimeConversionError):
                utc_to_et("2026-99-99T99:99:99")

    # Point 5: Target and observer verification
    def test_05_target_and_observer_verification(self):
        if not self.has_real_kernels:
            self.skipTest("Real kernels unavailable for target verification.")
        with SpiceManager(self.kernels_dir) as sm:
            sm.load_kernels([self.lsk, self.spk])

            # Supported target: Mars Barycenter (4)
            tid, name, ivs = sm.check_target_coverage(self.spk, "4")
            self.assertEqual(tid, 4)
            self.assertEqual(name, "MARS BARYCENTER")

            # Unsupported target
            with self.assertRaises(UnsupportedTargetError):
                sm.check_target_coverage(self.spk, 999999)

            # Identical target and observer rejection
            with self.assertRaises(InvalidParameterError):
                validate_extraction_parameters("4", "4", "J2000", "NONE")

    # Point 6: Coverage checks
    def test_06_coverage_checks(self):
        if not self.has_real_kernels:
            self.skipTest("Real kernels unavailable for coverage test.")
        with SpiceManager(self.kernels_dir) as sm:
            sm.load_kernel(self.lsk)
            # Epoch within DE440S (2026 is inside 1850-2150)
            et_valid = utc_to_et("2026-10-11T12:00:00")
            tid, name, ivs = sm.check_target_coverage(self.spk, 4, et=et_valid)
            self.assertEqual(tid, 4)

            # Epoch outside DE440S (year 3500)
            et_invalid = utc_to_et("3500-01-01T00:00:00")
            with self.assertRaises(CoverageError):
                sm.check_target_coverage(self.spk, 4, et=et_invalid)

    # Point 7: State vector extraction
    def test_07_state_vector_extraction(self):
        if not self.has_real_kernels:
            self.skipTest("Real kernels unavailable for state vector test.")
        with SpiceManager(self.kernels_dir) as sm:
            sm.load_kernels([self.lsk, self.spk])
            state = get_state_vector(
                target="4",
                observer="399",
                utc_timestamp="2026-10-11T00:00:00",
                reference_frame="J2000",
                aberration_correction="LT+S",
            )
            self.assertIsInstance(state, StateVector)
            # Verify coordinates are non-zero floats
            self.assertNotEqual(state.x_km, 0.0)
            self.assertNotEqual(state.vx_km_s, 0.0)
            self.assertGreater(state.light_time_s, 0.0)

            # Distance between Mars and Earth is physically ~2.38e8 km at this epoch
            dist = np.sqrt(state.x_km**2 + state.y_km**2 + state.z_km**2)
            self.assertGreater(dist, 2.0e8)
            self.assertLess(dist, 3.0e8)

    # Point 8: Time-series sampling
    def test_08_time_series_sampling(self):
        if not self.has_real_kernels:
            self.skipTest("Real kernels unavailable for time-series test.")
        with SpiceManager(self.kernels_dir) as sm:
            sm.load_kernels([self.lsk, self.spk])
            res = extract_time_series(
                target="4",
                observer="399",
                start_utc="2026-10-11T00:00:00",
                end_utc="2026-10-11T06:00:00",
                step_hours=2.0,
                reference_frame="J2000",
                aberration_correction="LT+S",
            )
            df = res.dataframe
            self.assertEqual(len(df), 4)  # 00:00, 02:00, 04:00, 06:00
            self.assertEqual(list(df.columns), DATAFRAME_COLUMNS)
            self.assertEqual(res.success_count, 4)
            self.assertEqual(res.failure_count, 0)

    # Point 9: CSV export
    def test_09_csv_export(self):
        with tempfile.TemporaryDirectory() as tmp_dir:
            csv_path = Path(tmp_dir) / "test_traj.csv"
            sample_df = pd.DataFrame([
                {
                    "UTC": "2026-10-11T00:00:00",
                    "ET_seconds": 844948869.182,
                    "Target": "4",
                    "Observer": "399",
                    "Frame": "J2000",
                    "X_km": -143655413.3,
                    "Y_km": 172712191.2,
                    "Z_km": 80276152.4,
                    "VX_km_s": -13.9558,
                    "VY_km_s": -24.4560,
                    "VZ_km_s": -9.9345,
                }
            ], columns=DATAFRAME_COLUMNS)

            out = export_to_csv(sample_df, csv_path)
            self.assertTrue(csv_path.exists())
            reloaded = pd.read_csv(csv_path)
            self.assertEqual(list(reloaded.columns), DATAFRAME_COLUMNS)
            self.assertEqual(len(reloaded), 1)

    # Point 10: Excel export
    def test_10_excel_export(self):
        with tempfile.TemporaryDirectory() as tmp_dir:
            xlsx_path = Path(tmp_dir) / "test_traj.xlsx"
            sample_df = pd.DataFrame([
                {
                    "UTC": "2026-10-11T00:00:00",
                    "ET_seconds": 844948869.182,
                    "Target": "4",
                    "Observer": "399",
                    "Frame": "J2000",
                    "X_km": -143655413.3,
                    "Y_km": 172712191.2,
                    "Z_km": 80276152.4,
                    "VX_km_s": -13.9558,
                    "VY_km_s": -24.4560,
                    "VZ_km_s": -9.9345,
                }
            ], columns=DATAFRAME_COLUMNS)

            out = export_to_excel(sample_df, xlsx_path, sheet_name="SPICE_Trajectory")
            self.assertTrue(xlsx_path.exists())

            wb = openpyxl.load_workbook(xlsx_path)
            self.assertIn("SPICE_Trajectory", wb.sheetnames)
            ws = wb["SPICE_Trajectory"]
            self.assertEqual(ws.freeze_panes, "A2")
            self.assertIsNotNone(ws.auto_filter.ref)
            self.assertEqual(ws.cell(row=2, column=6).number_format, "#,##0.00")
            self.assertEqual(ws.cell(row=2, column=9).number_format, "0.0000")

    # Point 11: Summary report generation
    def test_11_summary_report_generation(self):
        sample_df = pd.DataFrame([
            {
                "UTC": "2026-10-11T00:00:00",
                "ET_seconds": 844948869.182,
                "Target": "4",
                "Observer": "399",
                "Frame": "J2000",
                "X_km": -143655413.3,
                "Y_km": 172712191.2,
                "Z_km": 80276152.4,
                "VX_km_s": -13.9558,
                "VY_km_s": -24.4560,
                "VZ_km_s": -9.9345,
            }
        ], columns=DATAFRAME_COLUMNS)

        cfg = {
            "target": "4",
            "target_name": "MARS BARYCENTER",
            "observer": "399",
            "observer_name": "EARTH",
            "reference_frame": "J2000",
            "aberration_correction": "LT+S",
            "start_utc": "2026-10-11T00:00:00",
            "end_utc": "2026-10-11T00:00:00",
            "sampling_interval_hours": 1.0,
            "kernel_directory": "kernels",
        }

        summary = generate_text_summary(sample_df, cfg)
        self.assertIn("MARS BARYCENTER", summary)
        self.assertIn("EARTH", summary)
        self.assertIn("Position Units:           kilometres (km)", summary)
        self.assertIn("Velocity Units:           kilometres per second (km/s)", summary)
        self.assertIn("[First Successful Record]", summary)
        self.assertIn("SCIENTIFIC DATA PROVENANCE NOTICE", summary)
        self.assertIn("mathematically", summary)

    # Point 12: Error handling and guaranteed kernel cleanup
    def test_12_error_handling_and_cleanup(self):
        self.assertEqual(spice.ktotal("ALL"), 0)

        # Faulty config with invalid target
        faulty_cfg = {
            "target": "INVALID_BODY_9999",
            "observer": "399",
            "reference_frame": "J2000",
            "aberration_correction": "LT+S",
            "start_utc": "2026-10-11T00:00:00",
            "end_utc": "2026-10-11T02:00:00",
            "sampling_interval_hours": 1.0,
            "output_formats": ["csv"],
            "kernel_directory": "kernels",
            "output_directory": "output",
        }

        success, result, err = run_pipeline_with_cleanup(faulty_cfg)
        self.assertFalse(success)
        self.assertIsNotNone(err)

        # Guaranteed that kernel pool is 100% clean after failure!
        self.assertEqual(spice.ktotal("ALL"), 0)


if __name__ == "__main__":
    unittest.main()
