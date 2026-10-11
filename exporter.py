"""
SPICE Trajectory Data Exporter Module.

Exports extracted SPICE ephemeris DataFrames to:
1. CSV: Clean RFC 4180 tabular format.
2. Excel: Professionally styled workbook with frozen header, filters, custom precision, and auto column widths.
3. Text Summary: Detailed diagnostic report distinguishing mathematical ephemeris integration
   from measured in-situ telemetry.
"""

from __future__ import annotations

import os
from pathlib import Path
from typing import Any, Dict, Optional, Union

import numpy as np
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ==============================================================================
# Overwrite Policy Handler
# ==============================================================================

def check_overwrite_policy(filepath: Path, policy: str = "overwrite") -> bool:
    """
    Check if a file can be written based on the overwrite policy.

    Policies:
    - 'overwrite': overwrite existing file.
    - 'error': raise FileExistsError if file already exists.
    - 'skip': do not write if file already exists, return False.

    Returns True if writing should proceed, False if skipped.
    """
    valid_policies = {"overwrite", "error", "skip"}
    norm_policy = policy.lower().strip()
    if norm_policy not in valid_policies:
        raise ValueError(
            f"Invalid overwrite policy '{policy}'. Must be one of: {', '.join(sorted(valid_policies))}."
        )

    if filepath.exists():
        if norm_policy == "error":
            raise FileExistsError(
                f"Output file '{filepath}' already exists and overwrite policy is 'error'."
            )
        elif norm_policy == "skip":
            return False

    return True


# ==============================================================================
# CSV Export
# ==============================================================================

def export_to_csv(
    df: pd.DataFrame,
    filepath: Union[str, Path],
    overwrite_policy: str = "overwrite",
) -> Optional[Path]:
    """
    Export DataFrame to CSV format.

    Creates parent directories automatically.
    """
    path = Path(filepath).resolve()
    path.parent.mkdir(parents=True, exist_ok=True)

    if not check_overwrite_policy(path, overwrite_policy):
        return None

    df.to_csv(path, index=False)
    return path


# ==============================================================================
# Excel Export
# ==============================================================================

def export_to_excel(
    df: pd.DataFrame,
    filepath: Union[str, Path],
    sheet_name: str = "SPICE_Trajectory",
    overwrite_policy: str = "overwrite",
    failures_df: Optional[pd.DataFrame] = None,
) -> Optional[Path]:
    """
    Export DataFrame to an Excel (.xlsx) workbook with advanced formatting:
    - Clear worksheet name ('SPICE_Trajectory')
    - Descriptive headers with styled header fill and font
    - Appropriate numerical precision (km: #,##0.00, km/s: 0.0000, ET: #,##0.000)
    - Readable, auto-fit column widths
    - Frozen header row ('A2')
    - Auto-filter applied to header
    - Optional second sheet for extraction failures
    """
    path = Path(filepath).resolve()
    path.parent.mkdir(parents=True, exist_ok=True)

    if not check_overwrite_policy(path, overwrite_policy):
        return None

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name

    # Styles
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")  # Navy Blue
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    center_align = Alignment(horizontal="center", vertical="center")
    right_align = Alignment(horizontal="right", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    # Append Header
    headers = list(df.columns)
    ws.append(headers)

    for col_idx, _ in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align

    # Append Data
    for row in df.itertuples(index=False):
        ws.append(list(row))

    # Apply Formats & Widths
    num_rows = ws.max_row
    for r in range(2, num_rows + 1):
        for c, col_name in enumerate(headers, start=1):
            cell = ws.cell(row=r, column=c)
            cell.font = data_font
            cell.border = thin_border

            if col_name == "UTC":
                cell.alignment = center_align
            elif col_name in {"Target", "Observer", "Frame"}:
                cell.alignment = center_align
            elif col_name == "ET_seconds":
                cell.alignment = right_align
                cell.number_format = "#,##0.000"
            elif col_name in {"X_km", "Y_km", "Z_km"}:
                cell.alignment = right_align
                cell.number_format = "#,##0.00"
            elif col_name in {"VX_km_s", "VY_km_s", "VZ_km_s"}:
                cell.alignment = right_align
                cell.number_format = "0.0000"

    # Freeze header row
    ws.freeze_panes = "A2"

    # Auto-filter
    ws.auto_filter.ref = ws.dimensions

    # Auto-fit column widths
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # Optional Failures Sheet
    if failures_df is not None and not failures_df.empty:
        ws_fail = wb.create_sheet(title="Extraction_Failures")
        fail_headers = list(failures_df.columns)
        ws_fail.append(fail_headers)

        fail_header_fill = PatternFill(start_color="C00000", end_color="C00000", fill_type="solid")  # Crimson
        for col_idx, _ in enumerate(fail_headers, start=1):
            cell = ws_fail.cell(row=1, column=col_idx)
            cell.fill = fail_header_fill
            cell.font = header_font
            cell.alignment = center_align

        for row in failures_df.itertuples(index=False):
            ws_fail.append(list(row))

        ws_fail.freeze_panes = "A2"
        ws_fail.auto_filter.ref = ws_fail.dimensions

        for col in ws_fail.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = max(len(str(cell.value or "")) for cell in col)
            ws_fail.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 60)

    wb.save(path)
    return path


# ==============================================================================
# Text Summary Report
# ==============================================================================

def generate_text_summary(
    df: pd.DataFrame,
    config: Dict[str, Any],
    failures_df: Optional[pd.DataFrame] = None,
) -> str:
    """
    Generate a detailed textual summary report adhering strictly to requirements:
    - Target and observer (with resolved names/IDs).
    - Frame and aberration correction.
    - Requested time interval and sampling interval.
    - Successful and failed sample counts.
    - Position and velocity units.
    - First and last successful records.
    - Minimum and maximum position components.
    - Explicit statement distinguishing kernel-derived ephemerides from measured observations.
    """
    valid_df = df.dropna(subset=["X_km", "Y_km", "Z_km", "VX_km_s", "VY_km_s", "VZ_km_s"])
    total_count = len(df)
    success_count = len(valid_df)
    fail_count = total_count - success_count if failures_df is None else len(failures_df)

    target_desc = f"{config.get('target', 'N/A')}"
    if config.get("target_name"):
        target_desc += f" ({config.get('target_name')})"

    obs_desc = f"{config.get('observer', 'N/A')}"
    if config.get("observer_name"):
        obs_desc += f" ({config.get('observer_name')})"

    lines = []
    lines.append("=" * 80)
    lines.append("SPICE TRAJECTORY EXTRACTION SUMMARY REPORT")
    lines.append("=" * 80)
    lines.append("")

    lines.append("1. CONFIGURATION & METADATA")
    lines.append("-" * 40)
    lines.append(f"  Target Body / Object:     {target_desc}")
    lines.append(f"  Observer Body / Center:   {obs_desc}")
    lines.append(f"  Reference Coordinate Frame: {config.get('reference_frame', 'J2000')}")
    lines.append(f"  Aberration Correction:    {config.get('aberration_correction', 'NONE')}")
    lines.append(f"  Start UTC:                {config.get('start_utc', 'N/A')}")
    lines.append(f"  End UTC:                  {config.get('end_utc', 'N/A')}")
    lines.append(f"  Sampling Interval:        {config.get('sampling_interval_hours', 'N/A')} hours")
    lines.append(f"  Kernel Directory:         {config.get('kernel_directory', 'kernels')}")
    lines.append("")

    lines.append("2. SAMPLE EXTRACTION STATISTICS")
    lines.append("-" * 40)
    lines.append(f"  Total Requested Samples:  {total_count}")
    lines.append(f"  Successful Samples:       {success_count}")
    lines.append(f"  Failed Samples:           {fail_count}")
    lines.append(f"  Position Units:           kilometres (km)")
    lines.append(f"  Velocity Units:           kilometres per second (km/s)")
    lines.append("")

    if not valid_df.empty:
        first_row = valid_df.iloc[0]
        last_row = valid_df.iloc[-1]

        # Calculate distances (norm of position vector)
        distances = np.sqrt(valid_df["X_km"]**2 + valid_df["Y_km"]**2 + valid_df["Z_km"]**2)
        speeds = np.sqrt(valid_df["VX_km_s"]**2 + valid_df["VY_km_s"]**2 + valid_df["VZ_km_s"]**2)

        lines.append("3. BOUNDARY RECORDS")
        lines.append("-" * 40)
        lines.append("  [First Successful Record]")
        lines.append(f"    UTC:         {first_row['UTC']}")
        lines.append(f"    ET (sec):    {first_row['ET_seconds']:,.3f}")
        lines.append(f"    Position:    X={first_row['X_km']:,.2f} km, Y={first_row['Y_km']:,.2f} km, Z={first_row['Z_km']:,.2f} km")
        lines.append(f"    Velocity:    VX={first_row['VX_km_s']:,.4f} km/s, VY={first_row['VY_km_s']:,.4f} km/s, VZ={first_row['VZ_km_s']:,.4f} km/s")
        lines.append(f"    Range:       {distances.iloc[0]:,.2f} km")
        lines.append("")

        lines.append("  [Last Successful Record]")
        lines.append(f"    UTC:         {last_row['UTC']}")
        lines.append(f"    ET (sec):    {last_row['ET_seconds']:,.3f}")
        lines.append(f"    Position:    X={last_row['X_km']:,.2f} km, Y={last_row['Y_km']:,.2f} km, Z={last_row['Z_km']:,.2f} km")
        lines.append(f"    Velocity:    VX={last_row['VX_km_s']:,.4f} km/s, VY={last_row['VY_km_s']:,.4f} km/s, VZ={last_row['VZ_km_s']:,.4f} km/s")
        lines.append(f"    Range:       {distances.iloc[-1]:,.2f} km")
        lines.append("")

        lines.append("4. POSITION & VELOCITY EXTREMA")
        lines.append("-" * 40)
        lines.append(f"  X Component (km):         Min = {valid_df['X_km'].min():,.2f}  |  Max = {valid_df['X_km'].max():,.2f}")
        lines.append(f"  Y Component (km):         Min = {valid_df['Y_km'].min():,.2f}  |  Max = {valid_df['Y_km'].max():,.2f}")
        lines.append(f"  Z Component (km):         Min = {valid_df['Z_km'].min():,.2f}  |  Max = {valid_df['Z_km'].max():,.2f}")
        lines.append(f"  Range Distance (km):      Min = {distances.min():,.2f}  |  Max = {distances.max():,.2f}")
        lines.append(f"  Relative Speed (km/s):    Min = {speeds.min():,.4f}  |  Max = {speeds.max():,.4f}")
        lines.append("")
    else:
        lines.append("3. BOUNDARY RECORDS & EXTREMA")
        lines.append("-" * 40)
        lines.append("  No successful records were retrieved for this interval.")
        lines.append("")

    lines.append("5. SCIENTIFIC DATA PROVENANCE NOTICE")
    lines.append("-" * 40)
    lines.append("  CRITICAL DISTINCTION:")
    lines.append("  The state vectors and trajectory parameters in this report are mathematically")
    lines.append("  derived calculations produced via numerical integration and planetary ephemeris")
    lines.append("  models loaded from SPICE kernels (e.g. NASA JPL DE440S, NAIF LSK, and PCK).")
    lines.append("  They represent modeled celestial kinematics and are NOT direct in-situ spacecraft")
    lines.append("  telemetry, optical tracking observations, or radar measurements.")
    lines.append("=" * 80)
    lines.append("")

    return "\n".join(lines)


def export_text_summary(
    summary_text: str,
    filepath: Union[str, Path],
    overwrite_policy: str = "overwrite",
) -> Optional[Path]:
    """
    Export text summary to a file.

    Creates parent directories automatically.
    """
    path = Path(filepath).resolve()
    path.parent.mkdir(parents=True, exist_ok=True)

    if not check_overwrite_policy(path, overwrite_policy):
        return None

    path.write_text(summary_text, encoding="utf-8")
    return path
